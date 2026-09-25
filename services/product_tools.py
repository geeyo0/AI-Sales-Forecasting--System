"""Product entry, business-owned units, and a review-before-save Excel importer."""
import json
import re
import secrets
import time
from pathlib import Path

import pymysql
from flask import abort, flash, g, redirect, render_template, request, session, url_for, send_from_directory
from .product_import_reader import read_excel, validate_products

PRODUCT_CATEGORIES = (
    'Bread', 'Pastry', 'Beverage', 'Meal', 'Snack', 'Dessert', 'Other',
)

CATEGORY_ICONS = {
    'Bread': '🍞', 'Pastry': '🥐', 'Beverage': '☕', 'Meal': '🍽️',
    'Snack': '🍪', 'Dessert': '🍰', 'Other': '📦',
}


def install_product_tools(app, get_db, business_required, validate_csrf):
    @app.route('/js/product-tools.js')
    def product_tools_script():
        return send_from_directory(Path(app.root_path) / "static" / "js", 'product_tools.js')

    def business():
        with get_db().cursor() as cursor:
            cursor.execute('SELECT business_id, business_name, business_type, created_at FROM businesses WHERE user_id=%s',
                           (g.user['user_id'],))
            result = cursor.fetchone()
        if result is None:
            abort(403)
        return result
    @app.route("/business/products/export")
    @business_required
    def export_products():
        from io import BytesIO
        from datetime import datetime
        from zoneinfo import ZoneInfo

        from flask import send_file
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        owner = business()
        search = request.args.get("search", "").strip()[:100]

        with get_db().cursor() as cursor:
            cursor.execute(
                """
                SELECT product_name, selling_unit, selling_price
                FROM products
                WHERE business_id = %s
                  AND LOCATE(%s, product_name) > 0
                ORDER BY product_name
                """,
                (owner["business_id"], search),
            )
            products = cursor.fetchall()

        if not products:
            flash("There are no products to export.", "product_error")
            return redirect(url_for("business_products", search=search))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Products"

        sheet.append(["Product name", "Selling unit", "Selling price"])

        for product in products:
            sheet.append([
                product["product_name"],
                product["selling_unit"],
                product["selling_price"],
            ])

            row_number = sheet.max_row

            # Save names and units as text, never as Excel formulas.
            sheet.cell(row_number, 1).data_type = "s"
            sheet.cell(row_number, 2).data_type = "s"
            sheet.cell(row_number, 3).number_format = '"₱"#,##0.00'

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3B5F49")

        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 22
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        export_date = datetime.now(
            ZoneInfo("Asia/Manila")
        ).strftime("%Y-%m-%d")

        response = send_file(
            output,
            as_attachment=True,
            download_name=f"foodcast-products-{export_date}.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        response.headers["Cache-Control"] = "no-store"
        return response

    def lock_business(db, business_id):
        with db.cursor() as cursor:
            cursor.execute('SELECT business_id FROM businesses WHERE business_id=%s FOR UPDATE',
                           (business_id,))
            cursor.fetchone()

    def units_for(business_id):
        db = get_db()
        try:
            db.begin()
            lock_business(db, business_id)
            with db.cursor() as cursor:
                cursor.execute('INSERT IGNORE INTO product_unit_setup (business_id) VALUES (%s)',
                               (business_id,))
                if cursor.rowcount:
                    for name in ('Piece', 'Serving', 'Bottle'):
                        cursor.execute('INSERT IGNORE INTO product_units (business_id,name) VALUES (%s,%s)',
                                       (business_id, name))
                cursor.execute('SELECT unit_id,name FROM product_units WHERE business_id=%s ORDER BY name',
                               (business_id,))
                result = cursor.fetchall()
            db.commit()
            return result
        except Exception:
            db.rollback()
            raise

    def save_rows(business_id, rows):
        db = get_db()
        try:
            db.begin()
            lock_business(db, business_id)
            with db.cursor() as cursor:
                cursor.execute('SELECT name FROM product_units WHERE business_id=%s', (business_id,))
                clean, errors = validate_products(rows, [unit['name'] for unit in cursor.fetchall()])
                if errors:
                    db.rollback()
                    return errors
                for row, (name, unit, price) in zip(rows, clean):
                    category = str(row.get('category', '')).strip()
                    if category not in PRODUCT_CATEGORIES:
                        category = 'Other'
                    cursor.execute('''INSERT INTO products
                        (business_id,product_name,selling_unit,selling_price,category) VALUES (%s,%s,%s,%s,%s)''',
                        (business_id, name, unit, price, category))
            db.commit()
            return []
        except pymysql.err.IntegrityError as error:
            db.rollback()
            if error.args[0] == 1062:
                return [f'Nothing was saved: "{name}" already exists or repeats another product. Change its name, or untick that row when importing.']
            app.logger.exception('Product save failed')
            return ['Nothing was saved. Please check your entries.']
        except pymysql.MySQLError:
            db.rollback()
            app.logger.exception('Product save failed')
            return ['Nothing was saved. Please try again.']

    @business_required
    def products_page():
        owner = business()
        units = units_for(owner['business_id'])
        values = {
            'product_name': '', 'selling_unit': '', 'selling_price': '', 'category': '',
        }
        errors = []

        if request.method == 'POST':
            validate_csrf()
            values = {key: request.form.get(key, '').strip() for key in values}
            errors = save_rows(owner['business_id'], [{
                'name': values['product_name'],
                'unit': values['selling_unit'],
                'price': values['selling_price'],
                'category': values['category'],
            }])
            if not errors:
                flash('Product saved.', 'product_success')
                return redirect(url_for('business_products'))

        with get_db().cursor() as cursor:
            cursor.execute('''
                SELECT
                    p.product_id,
                    p.product_name,
                    p.category,
                    p.selling_unit,
                    p.selling_price,
                    COALESCE(
                        SUM(
                            CASE
                                WHEN m.movement_type IN ('Stock In', 'Adjustment In')
                                THEN m.quantity
                                ELSE -m.quantity
                            END
                        ), 0
                    ) AS available
                FROM products AS p
                LEFT JOIN stock_movements AS m ON m.product_id = p.product_id
                WHERE p.business_id = %s
                GROUP BY p.product_id, p.product_name, p.category, p.selling_unit, p.selling_price
                ORDER BY p.product_name
            ''', (owner['business_id'],))
            products = cursor.fetchall()

        for product in products:
            available = product['available']
            if available <= 0:
                product['stock_status'] = 'Out of Stock'
                product['status_class'] = 'stock-status-error'
            elif available <= 5:
                product['stock_status'] = 'Low Stock'
                product['status_class'] = 'stock-status-empty'
            else:
                product['stock_status'] = 'In Stock'
                product['status_class'] = 'stock-status-available'

        stock_summary = {
            'total': len(products),
            'in_stock': sum(1 for p in products if p['stock_status'] == 'In Stock'),
            'low_stock': sum(1 for p in products if p['stock_status'] == 'Low Stock'),
            'out_of_stock': sum(1 for p in products if p['stock_status'] == 'Out of Stock'),
        }

        return render_template(
            'business/products_tools.html',
            business=owner,
            units=units,
            categories=PRODUCT_CATEGORIES,
            category_icons=CATEGORY_ICONS,
            products=products,
            stock_summary=stock_summary,
            values=values,
            errors=errors,
        )

    # Preserve the URL and endpoint used by all existing links.
    app.view_functions['business_products'] = products_page

    @app.route('/business/product-units', methods=['POST'])
    @business_required
    def manage_product_units():
        validate_csrf()
        owner = business()
        units_for(owner['business_id'])
        name = request.form.get('name', '').strip()
        action = request.form.get('action')
        db = get_db()
        try:
            db.begin()
            lock_business(db, owner['business_id'])
            with db.cursor() as cursor:
                if action == 'add':
                    if not 1 <= len(name) <= 40 or any(ord(c) < 32 for c in name):
                        raise ValueError('Enter a unit name of 1–40 characters.')
                    cursor.execute('INSERT INTO product_units (business_id,name) VALUES (%s,%s)',
                                   (owner['business_id'], name))
                    message = f'Unit "{name}" added.'
                elif action == 'delete':
                    cursor.execute('SELECT name FROM product_units WHERE business_id=%s AND unit_id=%s',
                                   (owner['business_id'], request.form.get('unit_id', '')))
                    unit = cursor.fetchone()
                    if not unit:
                        raise ValueError('This unit is no longer available.')
                    cursor.execute('SELECT product_id FROM products WHERE business_id=%s AND selling_unit=%s LIMIT 1',
                                   (owner['business_id'], unit['name']))
                    if cursor.fetchone():
                        raise ValueError('This unit is used by a product and cannot be deleted. Existing sales and stock must keep their original unit.')
                    cursor.execute('DELETE FROM product_units WHERE business_id=%s AND unit_id=%s',
                                   (owner['business_id'], request.form['unit_id']))
                    message = f'Unit "{unit["name"]}" deleted.'
                else:
                    raise ValueError('Choose Add or Delete.')
            db.commit()
            flash(message, 'product_success')
        except ValueError as error:
            db.rollback()
            flash(str(error), 'product_error')
        except pymysql.MySQLError as error:
            db.rollback()
            if error.args[0] == 1062:
                flash('That unit already exists.', 'product_error')
            else:
                app.logger.exception('Unit update failed')
                flash('The unit could not be updated. Try again.', 'product_error')
        destination = 'product_import' if request.form.get('return_to') == 'import' else 'business_products'
        return redirect(url_for(destination))

    draft_dir = Path(app.root_path) / 'instance' / 'product-imports'

    def draft_path():
        token = session.get('product_import_token', '')
        if not re.fullmatch(r'[0-9a-f]{32}', token):
            return None
        return draft_dir / f'{g.user["user_id"]}-{token}.json'

    def discard_draft():
        path = draft_path()
        if path:
            path.unlink(missing_ok=True)
        session.pop('product_import_token', None)

    def read_draft(owner):
        path = draft_path()
        if not path or not path.exists():
            return None
        draft = json.loads(path.read_text(encoding='utf-8'))
        if draft['business_id'] != owner['business_id'] or time.time() - draft['created'] > 3600:
            discard_draft()
            return None
        return draft

    def write_draft(draft):
        draft_dir.mkdir(parents=True, exist_ok=True)
        path = draft_path()
        temporary = path.with_suffix('.tmp')
        temporary.write_text(json.dumps(draft), encoding='utf-8')
        temporary.replace(path)

    @app.route('/business/products/<int:product_id>/delete', methods=['POST'])
    @business_required
    def delete_product(product_id):
        validate_csrf()
        owner = business()
        db = get_db()
        try:
            with db.cursor() as cursor:
                cursor.execute(
                    'SELECT product_id FROM products WHERE product_id=%s AND business_id=%s',
                    (product_id, owner['business_id']),
                )
                if cursor.fetchone() is None:
                    abort(404)
                cursor.execute('DELETE FROM products WHERE product_id=%s', (product_id,))
            flash('Product deleted.', 'product_success')
        except pymysql.err.IntegrityError as error:
            if error.args[0] == 1451:
                flash(
                    'This product has recorded sales or stock history and cannot be deleted.',
                    'product_error',
                )
            else:
                app.logger.exception('Product delete failed')
                flash('The product could not be deleted.', 'product_error')
        except pymysql.MySQLError:
            app.logger.exception('Database error deleting a product.')
            flash('The product could not be deleted. Try again.', 'product_error')
        return redirect(url_for("business_products"))

    @app.route(
        "/business/products/bulk-delete",
        methods=["POST"],
    )
    @business_required
    def bulk_delete_products():
        validate_csrf()
        owner = business()

        try:
            product_ids = list(
                dict.fromkeys(
                    int(value)
                    for value in request.form.getlist("product_id")
                )
            )
        except (TypeError, ValueError):
            product_ids = []

        if not product_ids:
            flash(
                "Select at least one product to delete.",
                "product_error",
            )
            return redirect(url_for("business_products"))

        database = get_db()
        placeholders = ", ".join(["%s"] * len(product_ids))

        try:
            database.begin()

            with database.cursor() as cursor:
                cursor.execute(
                    f"""
                    SELECT product_id
                    FROM products
                    WHERE business_id = %s
                      AND product_id IN ({placeholders})
                    FOR UPDATE
                    """,
                    (owner["business_id"], *product_ids),
                )

                existing_ids = {
                    row["product_id"]
                    for row in cursor.fetchall()
                }

                if existing_ids != set(product_ids):
                    raise ValueError(
                        "One or more selected products are unavailable."
                    )

                cursor.execute(
                    f"""
                    DELETE FROM products
                    WHERE business_id = %s
                      AND product_id IN ({placeholders})
                    """,
                    (owner["business_id"], *product_ids),
                )

            database.commit()

            flash(
                f"{len(product_ids)} selected product(s) deleted.",
                "product_success",
            )

        except ValueError as error:
            database.rollback()
            flash(str(error), "product_error")

        except pymysql.err.IntegrityError as error:
            database.rollback()

            if error.args[0] == 1451:
                flash(
                    "One or more selected products have sales or "
                    "stock history and cannot be deleted.",
                    "product_error",
                )
            else:
                app.logger.exception(
                    "Bulk product deletion failed."
                )
                flash(
                    "The selected products could not be deleted.",
                    "product_error",
                )

        except pymysql.MySQLError:
            database.rollback()
            app.logger.exception(
                "Database error deleting selected products."
            )
            flash(
                "The selected products could not be deleted.",
                "product_error",
            )

        return redirect(url_for("business_products"))

    @app.route('/business/products/import-review', methods=['GET', 'POST'])
    @business_required
    def product_import():
        # Also caps form submissions with many editable rows.
        request.max_content_length = 6 * 1024 * 1024
        owner = business()
        units = units_for(owner['business_id'])
        draft = read_draft(owner)
        errors = []
        if request.method == 'POST':
            validate_csrf()
            action = request.form.get('action')
            try:
                if action == 'cancel':
                    discard_draft()
                    return redirect(url_for('business_products'))
                if action == 'upload':
                    file = request.files.get('file')
                    if not file or Path(file.filename or '').suffix.lower() != '.xlsx':
                        raise ValueError('Choose an Excel .xlsx file.')
                    raw = read_excel(file.stream.read(5 * 1024 * 1024 + 1))
                    discard_draft()
                    session['product_import_token'] = secrets.token_hex(16)
                    draft = dict(raw, business_id=owner['business_id'], created=time.time(), stage='map')
                    write_draft(draft)
                elif draft is None:
                    raise ValueError('Your preview expired. Upload the file again.')
                elif action == 'back':
                    draft['stage'] = 'map'
                    write_draft(draft)
                elif action == 'map':
                    start = int(request.form.get('first_row', '1')) - 1
                    has_header = request.form.get('has_header') == 'yes'
                    width = len(draft['rows'][0])
                    name_col = int(request.form.get('name_col', '-1'))
                    price_col = int(request.form.get('price_col', '-1'))
                    unit_col = int(request.form.get('unit_col', '-1'))
                    selected = [name_col, price_col] + ([unit_col] if unit_col != -1 else [])
                    if (not 0 <= start < len(draft['rows']) or start > 19
                            or any(not 0 <= c < width for c in selected)
                            or len(selected) != len(set(selected))):
                        raise ValueError('Choose different columns for name, price, and unit, and a valid starting row.')
                    preview = []
                    first = start + int(has_header)
                    for index, row in enumerate(draft['rows'][first:], first + 1):
                        if not any(cell['value'] for cell in row):
                            continue
                        if any(row[c]['invalid'] for c in selected):
                            raise ValueError(f'Excel row {index}: a selected column contains a formula or error. Paste values in Excel, save, and upload again.')
                        preview.append({'source': index, 'name': row[name_col]['value'],
                            'price': row[price_col]['value'],
                            'unit': (row[unit_col]['value'] if unit_col != -1 else '')
                                    or request.form.get('default_unit', ''), 'include': True})
                    if not preview or len(preview) > 500:
                        raise ValueError('Choose a range containing 1–500 products.')
                    # Normalize recognized units, but leave unknown values visible for correction.
                    clean, errors = validate_products(preview, [u['name'] for u in units])
                    for row, (_, unit, _) in zip(preview, clean):
                        if unit:
                            row['unit'] = unit
                    draft.update(stage='review', preview=preview)
                    write_draft(draft)
                elif action == 'save' and draft['stage'] == 'review':
                    for i, row in enumerate(draft['preview']):
                        row.update(name=request.form.get(f'name_{i}', '').strip(),
                                   price=request.form.get(f'price_{i}', '').strip(),
                                   unit=request.form.get(f'unit_{i}', '').strip(),
                                   include=request.form.get(f'include_{i}') == 'yes')
                    write_draft(draft)
                    rows = [row for row in draft['preview'] if row['include']]
                    errors = save_rows(owner['business_id'], rows)
                    if not errors:
                        discard_draft()
                        flash(f'{len(rows)} products imported.', 'product_success')
                        return redirect(url_for('business_products'))
                else:
                    raise ValueError('Please follow the import steps in order.')
            except (ValueError, TypeError) as error:
                errors = [str(error) or 'Please check your choices.']
        if draft and draft['stage'] == 'review' and not errors:
            _, errors = validate_products([r for r in draft['preview'] if r['include']], [u['name'] for u in units])
        return render_template('business/import_products.html', business=owner, units=units,
                               draft=draft, errors=errors[:20])
